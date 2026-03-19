from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = "/Users/leo/Desktop/open-project/zsad/visual_prompt_architecture.pptx"


def add_box(slide, left, top, width, height, text, fill, line=(60, 60, 60), font_size=18, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(*line)
    text_frame = shape.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(20, 20, 20)
    return shape


def add_text(slide, left, top, width, height, text, font_size=16, bold=False, color=(30, 30, 30)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def connect(slide, shape_a, shape_b, color=(90, 90, 90), begin_side="right", end_side="left"):
    ax = shape_a.left + (shape_a.width if begin_side == "right" else shape_a.width // 2)
    ay = shape_a.top + shape_a.height // 2
    bx = shape_b.left if end_side == "left" else shape_b.left + shape_b.width // 2
    by = shape_b.top + shape_b.height // 2
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, ay, bx, by)
    line.line.color.rgb = RGBColor(*color)
    line.line.width = Pt(1.5)
    return line


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, Inches(0.45), Inches(0.2), Inches(5), Inches(0.5), "视觉 Prompt 融合后的 AnomalyCLIP 结构", 24, True)
    add_text(slide, Inches(0.45), Inches(0.65), Inches(7), Inches(0.35), "基于当前 train_visual_prompt.py 的实际实现", 11, False, (90, 90, 90))

    image_in = add_box(slide, Inches(0.4), Inches(1.2), Inches(1.4), Inches(0.55), "输入图像", (223, 240, 216), font_size=20, bold=True)
    text_in = add_box(slide, Inches(0.4), Inches(2.15), Inches(1.4), Inches(0.55), "文本 Prompt", (222, 235, 247), font_size=20, bold=True)

    visual_encoder = add_box(slide, Inches(2.15), Inches(1.0), Inches(1.8), Inches(0.9), "AnomalyCLIP\n视觉编码器", (253, 233, 217), font_size=20, bold=True)
    text_encoder = add_box(slide, Inches(2.15), Inches(2.0), Inches(1.8), Inches(0.9), "PromptLearner +\n文本编码器", (242, 242, 242), font_size=20, bold=True)

    global_v = add_box(slide, Inches(4.45), Inches(0.95), Inches(1.55), Inches(0.6), "全局视觉特征 V", (252, 228, 214), font_size=18)
    patch_v = add_box(slide, Inches(4.45), Inches(1.8), Inches(1.55), Inches(0.6), "多层 Patch 特征", (252, 228, 214), font_size=18)
    text_t = add_box(slide, Inches(4.45), Inches(2.65), Inches(1.55), Inches(0.6), "文本特征 T", (224, 224, 224), font_size=18)

    cross_attn = add_box(slide, Inches(6.35), Inches(1.25), Inches(1.65), Inches(0.7), "Cross-Attention\nQ=V, K/V=T", (198, 224, 180), font_size=18, bold=True)
    vp = add_box(slide, Inches(8.25), Inches(1.25), Inches(1.35), Inches(0.7), "视觉 Prompt\nV~", (255, 242, 204), font_size=18, bold=True)
    mlp = add_box(slide, Inches(6.35), Inches(2.4), Inches(1.65), Inches(0.7), "MLP + Residual\n+ Norm", (255, 217, 102), font_size=18, bold=True)
    enhanced = add_box(slide, Inches(8.25), Inches(2.4), Inches(1.35), Inches(0.7), "增强特征 V'", (244, 204, 204), font_size=18, bold=True)

    sim = add_box(slide, Inches(6.2), Inches(3.65), Inches(1.8), Inches(0.7), "Patch-Text\nSimilarity", (221, 235, 247), font_size=18, bold=True)
    anomaly_map = add_box(slide, Inches(8.35), Inches(3.65), Inches(1.55), Inches(0.7), "Anomaly Maps", (234, 209, 220), font_size=18)
    hsf = add_box(slide, Inches(10.2), Inches(3.65), Inches(1.25), Inches(0.7), "HSF", (217, 234, 211), font_size=20, bold=True)
    fusion = add_box(slide, Inches(10.0), Inches(2.3), Inches(1.7), Inches(0.8), "融合\n0.8·V' + 0.2·Fc", (208, 224, 227), font_size=18, bold=True)
    output = add_box(slide, Inches(11.95), Inches(2.3), Inches(1.0), Inches(0.8), "最终判别", (226, 239, 218), font_size=18, bold=True)

    connect(slide, image_in, visual_encoder)
    connect(slide, text_in, text_encoder)
    connect(slide, visual_encoder, global_v)
    connect(slide, visual_encoder, patch_v)
    connect(slide, text_encoder, text_t)
    connect(slide, global_v, cross_attn)
    connect(slide, text_t, cross_attn)
    connect(slide, cross_attn, vp)
    connect(slide, vp, mlp)
    connect(slide, global_v, mlp)
    connect(slide, mlp, enhanced)
    connect(slide, patch_v, sim)
    connect(slide, text_t, sim)
    connect(slide, sim, anomaly_map)
    connect(slide, anomaly_map, hsf)
    connect(slide, hsf, fusion)
    connect(slide, enhanced, fusion, begin_side="right", end_side="left")
    connect(slide, fusion, output)

    add_text(slide, Inches(10.0), Inches(4.65), Inches(2.7), Inches(1.0), "Fc: HSF 从高异常 patch 中聚合得到的聚类特征", 13, False, (70, 70, 70))
    add_text(slide, Inches(0.55), Inches(5.75), Inches(12.0), Inches(0.7), "核心改动: 用文本特征指导生成视觉 Prompt，再把它和原始视觉特征融合，得到更贴近异常语义的图像表示 V'。", 16, False)

    # Slide 2
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide2, Inches(0.45), Inches(0.2), Inches(6), Inches(0.5), "训练目标与实现映射", 24, True)

    add_box(slide2, Inches(0.5), Inches(1.0), Inches(3.5), Inches(1.1), "公式 1\nV~ = softmax(VTᵀ / √D) · T", (255, 242, 204), font_size=24, bold=True)
    add_box(slide2, Inches(0.5), Inches(2.4), Inches(3.5), Inches(1.1), "公式 2\nV' = MLP([V ; V~])", (244, 204, 204), font_size=24, bold=True)

    add_box(slide2, Inches(4.4), Inches(0.95), Inches(3.9), Inches(0.8), "图像级分类损失: CrossEntropy(V', T)", (226, 239, 218), font_size=20, bold=True)
    add_box(slide2, Inches(4.4), Inches(1.95), Inches(3.9), Inches(0.8), "异常分类损失: FocalLoss(融合特征, 标签)", (217, 234, 211), font_size=20, bold=True)
    add_box(slide2, Inches(4.4), Inches(2.95), Inches(3.9), Inches(0.8), "像素级分割损失: Focal + Dice", (221, 235, 247), font_size=20, bold=True)
    add_box(slide2, Inches(4.4), Inches(3.95), Inches(3.9), Inches(0.8), "KD 损失: LLM Embedding 与 Prompt 对齐", (234, 209, 220), font_size=20, bold=True)
    add_box(slide2, Inches(4.4), Inches(4.95), Inches(3.9), Inches(0.8), "视觉 Prompt 对齐损失: MSE(V', V~)", (208, 224, 227), font_size=20, bold=True)

    add_box(slide2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(2.6), "代码对应\n- visual_prompt.py: VisualPromptAligner\n- train_visual_prompt.py: 接入训练\n- checkpoint: 同时保存 prompt_learner + visual_prompt", (242, 242, 242), font_size=18)
    add_text(slide2, Inches(0.65), Inches(6.1), Inches(12.0), Inches(0.7), "建议汇报表述: 视觉 Prompt 不是替换 CLIP 主干，而是在冻结视觉编码器的前提下，用文本语义对视觉表示做轻量对齐增强。", 16, False)

    # Slide 3
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide3, Inches(0.45), Inches(0.2), Inches(6.2), Inches(0.5), "小创新点: 自适应双门控融合", 24, True)
    add_text(slide3, Inches(0.45), Inches(0.62), Inches(8), Inches(0.35), "目的: 避免固定融合系数，按样本动态控制文本引导信息和聚类异常信息的注入强度", 11, False, (90, 90, 90))

    gv = add_box(slide3, Inches(0.55), Inches(1.4), Inches(1.55), Inches(0.7), "原始视觉特征 V", (252, 228, 214), font_size=18, bold=True)
    vpv = add_box(slide3, Inches(0.55), Inches(2.5), Inches(1.55), Inches(0.7), "视觉 Prompt 特征", (255, 242, 204), font_size=18, bold=True)
    gate1 = add_box(slide3, Inches(2.45), Inches(1.82), Inches(1.5), Inches(0.8), "Gate 1\n自适应视觉融合", (217, 234, 211), font_size=18, bold=True)
    gvf = add_box(slide3, Inches(4.25), Inches(1.82), Inches(1.7), Inches(0.8), "门控视觉特征\nVg", (244, 204, 204), font_size=18, bold=True)

    fc = add_box(slide3, Inches(6.4), Inches(2.5), Inches(1.55), Inches(0.7), "HSF 聚类特征 Fc", (208, 224, 227), font_size=18, bold=True)
    gate2 = add_box(slide3, Inches(6.35), Inches(1.05), Inches(1.6), Inches(0.8), "Gate 2\n自适应聚类融合", (217, 234, 211), font_size=18, bold=True)
    fout = add_box(slide3, Inches(8.35), Inches(1.45), Inches(1.8), Inches(0.8), "最终图像特征\nVfinal", (226, 239, 218), font_size=18, bold=True)
    pred = add_box(slide3, Inches(10.55), Inches(1.45), Inches(1.8), Inches(0.8), "异常分类 / 检测", (242, 242, 242), font_size=18, bold=True)

    connect(slide3, gv, gate1)
    connect(slide3, vpv, gate1)
    connect(slide3, gate1, gvf)
    connect(slide3, gvf, gate2)
    connect(slide3, fc, gate2)
    connect(slide3, gate2, fout)
    connect(slide3, fout, pred)

    add_box(slide3, Inches(0.65), Inches(4.25), Inches(5.6), Inches(1.15), "公式 1\nGv = sigmoid(MLP([V ; V']))\nVg = Gv ⊙ V' + (1 - Gv) ⊙ V", (255, 242, 204), font_size=22, bold=True)
    add_box(slide3, Inches(6.55), Inches(4.25), Inches(5.9), Inches(1.15), "公式 2\nGc = sigmoid(MLP([Vg ; Fc]))\nVfinal = Gc ⊙ Fc + (1 - Gc) ⊙ Vg", (244, 204, 204), font_size=22, bold=True)
    add_text(slide3, Inches(0.7), Inches(5.8), Inches(12.0), Inches(0.8), "实现位置: adaptive_fusion.py + train_visual_prompt_gate.py。额外加入 gate regularization，抑制门控过早塌缩到极端值。", 16, False)

    prs.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
